"""Build editable Word exports for generated reading exercises."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def _text(value: object, name: str, limit: int = 100000) -> str:
    if not isinstance(value, str) or not value.strip() or len(value) > limit:
        raise ValueError(f"{name}内容无效。")
    return value.strip()


def _east_asian_font(style, name: str) -> None:
    fonts = style.element.get_or_add_rPr().rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        style.element.get_or_add_rPr().insert(0, fonts)
    fonts.set(qn("w:eastAsia"), name)


def _add_paragraphs(document: Document, value: str) -> None:
    for block in value.replace("\r\n", "\n").split("\n\n"):
        text = block.strip()
        if text:
            document.add_paragraph(text)


def build_reading_docx(payload: dict) -> bytes:
    if not isinstance(payload, dict):
        raise ValueError("导出内容格式无效。")
    kind = payload.get("kind")
    if kind not in {"story", "exam"}:
        raise ValueError("不支持这个阅读导出类型。")
    title = _text(payload.get("title"), "标题", 200)
    passage = _text(payload.get("passage"), "英文原文")
    translation = _text(payload.get("translation"), "中文译文")

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.86)
    section.right_margin = Inches(0.86)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(38, 55, 47)
    _east_asian_font(normal, "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.4
    for name, size in (("Title", 20), ("Heading 1", 13), ("Heading 2", 11)):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(35, 85, 55)
        _east_asian_font(style, "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(7)

    document.add_paragraph(title, style="Title")
    if kind == "exam":
        exam = _text(payload.get("exam"), "考试类型", 80)
        document.add_paragraph(f"{exam} · 原创模拟阅读")
    else:
        words = payload.get("words", [])
        if not isinstance(words, list) or len(words) > 500 or not all(isinstance(word, str) and len(word) <= 80 for word in words):
            raise ValueError("目标单词格式无效。")
        document.add_paragraph("阅读短文练习")

    document.add_heading("英文原文", level=1)
    _add_paragraphs(document, passage)

    if kind == "exam":
        questions = payload.get("questions")
        if not isinstance(questions, list) or len(questions) != 5:
            raise ValueError("阅读题目格式无效。")
        document.add_heading("阅读理解", level=1)
        answers = []
        for index, question in enumerate(questions, 1):
            if not isinstance(question, dict):
                raise ValueError("阅读题目格式无效。")
            stem = _text(question.get("stem"), "题干", 2000)
            options = question.get("options")
            if not isinstance(options, list) or len(options) != 4:
                raise ValueError("选项格式无效。")
            answer = _text(question.get("answer"), "答案", 1)
            if answer not in "ABCD":
                raise ValueError("答案格式无效。")
            explanation = _text(question.get("explanation"), "解析", 3000)
            document.add_heading(f"{index}. {stem}", level=2)
            for letter, option in zip("ABCD", options):
                document.add_paragraph(f"{letter}. {_text(option, '选项', 2000)}")
            answers.append((answer, explanation))
        document.add_heading("答案与解析", level=1)
        for index, (answer, explanation) in enumerate(answers, 1):
            document.add_paragraph(f"{index}. {answer}  {explanation}")

    document.add_heading("中文译文", level=1)
    _add_paragraphs(document, translation)
    if kind == "story" and words:
        document.add_heading("目标单词", level=1)
        document.add_paragraph("、".join(words))

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("词页 · 英语练习工作台").font.size = Pt(8)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
