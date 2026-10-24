import io
import unittest

from docx import Document

from export_docx import build_reading_docx


class ReadingWordExportTests(unittest.TestCase):
    def test_story_contains_original_translation_and_words(self):
        result = build_reading_docx({"kind": "story", "title": "A Journey", "passage": "A curious journey begins.",
                                     "translation": "一次好奇的旅程开始了。", "words": ["curious", "journey"]})
        self.assertTrue(result.startswith(b"PK"))
        paragraphs = [paragraph.text for paragraph in Document(io.BytesIO(result)).paragraphs]
        self.assertIn("A curious journey begins.", paragraphs)
        self.assertIn("一次好奇的旅程开始了。", paragraphs)
        self.assertIn("curious、journey", paragraphs)

    def test_exam_contains_questions_answers_and_translation(self):
        questions = [{"stem": f"Question {index}?", "options": ["One", "Two", "Three", "Four"],
                      "answer": "A", "explanation": "答案来自文章。"} for index in range(1, 6)]
        result = build_reading_docx({"kind": "exam", "title": "Practice", "exam": "英语四级", "passage": "A short passage.",
                                     "translation": "一篇短文。", "questions": questions})
        paragraphs = [paragraph.text for paragraph in Document(io.BytesIO(result)).paragraphs]
        self.assertIn("英语四级 · 原创模拟阅读", paragraphs)
        self.assertIn("1. Question 1?", paragraphs)
        self.assertIn("A. One", paragraphs)
        self.assertIn("1. A  答案来自文章。", paragraphs)
        self.assertIn("一篇短文。", paragraphs)
        self.assertNotIn("_export_answer", questions[0])

    def test_invalid_exam_is_rejected(self):
        with self.assertRaises(ValueError):
            build_reading_docx({"kind": "exam", "title": "X", "exam": "英语四级", "passage": "Text.",
                                "translation": "译文。", "questions": []})


if __name__ == "__main__":
    unittest.main()
